import unicodedata
from io import BytesIO

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.errores import DatosInvalidos
from app.models.categoria_model import Categoria
from app.models.producto_model import Producto
from app.schemas.importacion_schema import FilaInventario, ResultadoImportacion
from app.services import categoria_service, producto_service

COLUMNAS = ("nombre", "precio", "stock", "categoria")
TAMANO_MAXIMO = 2 * 1024 * 1024
FILAS_MAXIMAS = 2000

MENSAJES = {
    "string_type": "debe ser texto",
    "string_too_short": "es muy corto",
    "string_too_long": "es muy largo",
    "int_type": "debe ser un número entero",
    "int_parsing": "debe ser un número entero",
    "int_from_float": "no puede tener decimales",
    "greater_than": "debe ser mayor que cero",
    "greater_than_equal": "no puede ser negativo",
}


def _normalizar(texto) -> str:
    sin_tildes = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return sin_tildes.strip().lower()


def _traducir(error: dict) -> str:
    campo = error["loc"][0]
    if error.get("input") is None:
        return f"{campo}: la celda está vacía"
    return f"{campo}: {MENSAJES.get(error['type'], error['msg'])} (llegó: {error['input']!r})"


def _leer_filas(nombre_archivo: str, contenido: bytes) -> list[tuple]:
    if not nombre_archivo.lower().endswith(".xlsx"):
        raise DatosInvalidos("Solo se aceptan archivos de Excel con extensión .xlsx")
    if len(contenido) > TAMANO_MAXIMO:
        raise DatosInvalidos("El archivo pesa más de 2 MB")
    try:
        libro = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    except Exception as error:
        raise DatosInvalidos("El archivo no se pudo abrir como Excel (.xlsx)") from error
    filas = list(libro.active.iter_rows(values_only=True))
    libro.close()
    if len(filas) < 2:
        raise DatosInvalidos("El archivo no tiene productos para importar")
    if len(filas) - 1 > FILAS_MAXIMAS:
        raise DatosInvalidos(f"El archivo supera el máximo de {FILAS_MAXIMAS} productos")
    return filas


def _ubicar_columnas(encabezado: tuple) -> dict[str, int]:
    nombres = [_normalizar(celda) if celda is not None else "" for celda in encabezado]
    faltantes = [columna for columna in COLUMNAS if columna not in nombres]
    if faltantes:
        raise DatosInvalidos(f"A la primera fila le faltan las columnas: {', '.join(faltantes)}")
    return {columna: nombres.index(columna) for columna in COLUMNAS}


def _validar_filas(filas: list[tuple], posicion: dict[str, int], nombres_existentes: set[str]):
    validas: list[FilaInventario] = []
    errores: list[str] = []
    for numero, fila in enumerate(filas, start=2):
        if all(celda is None or str(celda).strip() == "" for celda in fila):
            continue
        datos = {columna: fila[i] if i < len(fila) else None for columna, i in posicion.items()}
        try:
            item = FilaInventario.model_validate(datos)
        except ValidationError as error:
            errores.extend(f"Fila {numero} · {_traducir(e)}" for e in error.errors())
            continue
        if _normalizar(item.nombre) in nombres_existentes:
            errores.append(f"Fila {numero} · nombre: '{item.nombre}' ya existe en el inventario o está repetido en el archivo")
            continue
        nombres_existentes.add(_normalizar(item.nombre))
        validas.append(item)
    return validas, errores


def _guardar(db: Session, validas: list[FilaInventario]) -> list[str]:
    categorias = {_normalizar(c.nombre): c for c in categoria_service.listar_categorias(db)}
    creadas: list[str] = []
    for item in validas:
        categoria = categorias.get(_normalizar(item.categoria))
        if categoria is None:
            categoria = Categoria(nombre=item.categoria)
            db.add(categoria)
            categorias[_normalizar(item.categoria)] = categoria
            creadas.append(item.categoria)
        db.add(Producto(nombre=item.nombre, precio=item.precio, stock=item.stock, categoria=categoria))
    db.commit()
    return creadas


def importar_inventario(db: Session, nombre_archivo: str, contenido: bytes) -> ResultadoImportacion:
    filas = _leer_filas(nombre_archivo, contenido)
    posicion = _ubicar_columnas(filas[0])
    nombres_existentes = {_normalizar(p.nombre) for p in producto_service.listar_productos(db)}
    validas, errores = _validar_filas(filas[1:], posicion, nombres_existentes)
    if errores:
        raise DatosInvalidos(f"El archivo tiene {len(errores)} error(es). No se importó ningún producto", errores)
    if not validas:
        raise DatosInvalidos("El archivo no tiene productos para importar")
    categorias_creadas = _guardar(db, validas)
    return ResultadoImportacion(productos_creados=len(validas), categorias_creadas=categorias_creadas)